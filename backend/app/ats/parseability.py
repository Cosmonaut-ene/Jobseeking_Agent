"""Parseability checker — SPEC.md 附录 F.5 TASK-C01.

Deterministic, no LLM call. Re-extracts text from the user's actual
uploaded resume file the same way ResumeParser would, then checks how
much of the profile's key fields survive that extraction. A real ATS
parser can fail the exact same way (tables, multi-column layout,
scanned images) regardless of how well-written the content is — this
module measures that risk directly instead of asking an LLM to guess it.
"""
from __future__ import annotations

import re
from pathlib import Path

from pydantic import BaseModel

from backend.app.agents.parser import extract_text_from_file
from backend.app.models.user_profile import UserProfile

_LOW_TEXT_CHAR_THRESHOLD = 50
_LOW_SCORE_WARNING_THRESHOLD = 70.0


class ParseabilityReport(BaseModel):
    score: float  # 0-100, 命中率
    extracted_char_count: int
    missing_fields: list[str]
    warnings: list[str]


def _expected_fields(profile: UserProfile) -> list[tuple[str, str]]:
    """姓名 / 各项技能名 / 公司名 / 学校名 —— 与 SPEC 附录 F.5 TASK-C01 一致。"""
    fields: list[tuple[str, str]] = []
    if profile.name:
        fields.append(("姓名", profile.name))
    for skill in profile.skills:
        if skill.name:
            fields.append((f"技能: {skill.name}", skill.name))
    for exp in profile.experience:
        if exp.company:
            fields.append((f"公司: {exp.company}", exp.company))
    for edu in profile.education:
        if edu.institution:
            fields.append((f"学校: {edu.institution}", edu.institution))
    return fields


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower()


def _has_docx_tables_with_content(file_path: Path) -> bool:
    """python-docx 的段落级抽取（`doc.paragraphs`）不读取表格单元格内容——
    简历如果用表格排版联系方式/技能栏，这部分内容会被现有抽取逻辑静默丢弃。
    """
    if file_path.suffix.lower() != ".docx":
        return False
    try:
        from docx import Document

        doc = Document(str(file_path))
        return any(
            cell.text.strip()
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        )
    except Exception:
        return False


def check_parseability(file_path: Path | str, expected_profile: UserProfile) -> ParseabilityReport:
    """核对同一份文件重新抽取出的文本，能覆盖 expected_profile 里多少关键字段。

    确定性：同一 (file_path, expected_profile) 输入永远得到同一份报告，
    不调用任何 LLM。文件损坏/无法解析时优雅降级为全丢失报告，不抛异常。
    """
    file_path = Path(file_path)
    fields = _expected_fields(expected_profile)

    try:
        raw_text = extract_text_from_file(file_path)
    except Exception as exc:
        return ParseabilityReport(
            score=0.0,
            extracted_char_count=0,
            missing_fields=[f"{label} — 未在提取文本中找到" for label, _ in fields],
            warnings=[f"文件解析失败，无法抽取任何文本：{exc}"],
        )

    normalized_text = _normalize(raw_text)
    missing = [
        f"{label} — 未在提取文本中找到"
        for label, value in fields
        if _normalize(value) not in normalized_text
    ]
    score = 100.0 if not fields else round((len(fields) - len(missing)) / len(fields) * 100, 1)

    char_count = len(raw_text.strip())
    warnings: list[str] = []
    if char_count < _LOW_TEXT_CHAR_THRESHOLD:
        warnings.append("几乎未提取到文本，文件可能是扫描件/纯图片，或格式不受支持")
    elif fields and score < _LOW_SCORE_WARNING_THRESHOLD:
        warnings.append("大量字段未在提取文本中命中，简历可能使用了表格/多栏/图文混排导致内容丢失")

    if _has_docx_tables_with_content(file_path):
        warnings.append("检测到 Word 表格；当前抽取逻辑只读取段落文本，表格内文字不计入解析结果")

    return ParseabilityReport(
        score=score,
        extracted_char_count=char_count,
        missing_fields=missing,
        warnings=warnings,
    )
