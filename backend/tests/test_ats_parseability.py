"""Tests for the deterministic ATS parseability checker (SPEC 附录 F.5 TASK-C01)."""
import pytest

from backend.app.ats.parseability import check_parseability
from backend.app.models.user_profile import Education, Experience, Preferences, Skill, UserProfile


def _make_profile() -> UserProfile:
    return UserProfile(
        name="Jane Doe",
        target_roles=["Software Engineer"],
        skills=[Skill(name="Kubernetes", level="intermediate", years=2.0)],
        experience=[Experience(company="Acme Corp", role="Engineer", duration="2022-2024")],
        projects=[],
        education=[Education(institution="Test University", degree="BSc")],
        preferences=Preferences(),
    )


def test_high_score_when_all_fields_present(tmp_path, monkeypatch):
    """All key fields findable in extracted text -> score 100, no missing fields."""
    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"dummy")  # content irrelevant, extraction is monkeypatched below
    monkeypatch.setattr(
        "backend.app.ats.parseability.extract_text_from_file",
        lambda path: (
            "Jane Doe\nSenior Software Engineer\n"
            "Skills: Kubernetes, Python, Docker\n"
            "Experience: Acme Corp — Engineer (2022-2024)\n"
            "Education: Test University — BSc Computer Science"
        ),
    )

    report = check_parseability(resume, _make_profile())

    assert report.score >= 90
    assert report.missing_fields == []
    assert report.warnings == []


def test_all_fields_missing_flags_low_score_and_warning(tmp_path, monkeypatch):
    """None of the expected fields appear -> score 0, warning raised."""
    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"dummy")
    monkeypatch.setattr(
        "backend.app.ats.parseability.extract_text_from_file",
        lambda path: "Some completely unrelated placeholder text that is long enough " * 3,
    )

    report = check_parseability(resume, _make_profile())

    assert report.score == 0.0
    assert len(report.missing_fields) == 4  # name + skill + company + school
    assert any("命中" in w for w in report.warnings)


def test_normal_docx_real_extraction(tmp_path):
    """Real python-docx round trip, no mocking — proves the actual pipeline works."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Experience: Acme Corp, Engineer")
    doc.add_paragraph("Skills: Kubernetes")
    doc.add_paragraph("Education: Test University")
    path = tmp_path / "resume.docx"
    doc.save(path)

    report = check_parseability(path, _make_profile())

    assert report.score >= 90
    assert report.missing_fields == []


def test_docx_table_content_not_extracted_flags_warning(tmp_path):
    """Key field placed only inside a table cell — paragraph-only extraction misses it,
    and the checker must flag this as a known structural risk, not silently pass."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("Skills: Kubernetes")
    doc.add_paragraph("Education: Test University")
    table = doc.add_table(rows=1, cols=1)
    table.rows[0].cells[0].text = "Acme Corp"  # company only in a table cell
    path = tmp_path / "resume.docx"
    doc.save(path)

    report = check_parseability(path, _make_profile())

    assert any("Acme Corp" in m for m in report.missing_fields)
    assert any("表格" in w for w in report.warnings)


def test_corrupted_file_returns_graceful_report_not_exception(tmp_path):
    """A file that fails to parse must degrade to a full-miss report, never raise."""
    resume = tmp_path / "resume.pdf"
    resume.write_bytes(b"\x00\x01not a real pdf at all")

    report = check_parseability(resume, _make_profile())

    assert report.score == 0.0
    assert report.extracted_char_count == 0
    assert len(report.missing_fields) == 4
    assert any("解析失败" in w for w in report.warnings)


def test_score_is_deterministic_across_repeated_calls(tmp_path):
    """Same (file, profile) input must always produce the same report."""
    from docx import Document

    doc = Document()
    doc.add_paragraph("Jane Doe — Kubernetes — Acme Corp — Test University")
    path = tmp_path / "resume.docx"
    doc.save(path)
    profile = _make_profile()

    reports = [check_parseability(path, profile) for _ in range(10)]

    assert len({r.score for r in reports}) == 1
    assert len({tuple(r.missing_fields) for r in reports}) == 1
